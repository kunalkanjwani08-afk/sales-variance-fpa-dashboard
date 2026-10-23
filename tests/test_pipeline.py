"""End-to-end tests: generate -> clean -> analyse -> insights -> dashboard, in a temp folder.

These are the slowest tests (about a minute) because they run the real pipeline
on the full 24-month dataset.
"""
import json
import re
from decimal import Decimal

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from fpa_lib import config
from tests.conftest import ROOT

pytestmark = pytest.mark.slow


def _read_clean(path):
    return pd.read_csv(path, dtype={"Line_ID": str, "Notes": str, "Repaired_Fields": str, "DQ_Flag": str},
                       keep_default_na=False, na_values=[""])


# --------------------------------------------------------------- run ok
def test_every_step_exits_zero(pipeline):
    assert pipeline["codes"] == {"gen": 0, "clean": 0, "analysis": 0, "insights": 0, "dashboard": 0}


# ----------------------------------------------------------- generator
def test_generator_is_deterministic(scripts):
    g = scripts["gen"]
    t1, t2 = g.build_truth(config.RANDOM_SEED), g.build_truth(config.RANDOM_SEED)
    pd.testing.assert_frame_equal(t1, t2)
    r1, m1 = g.corrupt(t1, config.RANDOM_SEED)
    r2, m2 = g.corrupt(t2, config.RANDOM_SEED)
    pd.testing.assert_frame_equal(r1, r2)
    pd.testing.assert_frame_equal(m1, m2)
    assert not g.build_truth(config.RANDOM_SEED + 1).equals(t1)


def test_generated_data_is_not_a_grid(scripts):
    """v1 had one price per product and one line per region-product-month. v2 must not."""
    t = scripts["gen"].build_truth(config.RANDOM_SEED)
    per_cell = t.groupby(["Region", "Product", "Date"]).size()
    assert per_cell.nunique() > 1                                     # varying number of lines per cell
    assert (t.groupby("Product")["Actual_Unit_Price_USD"].nunique() > 10).all()
    assert (t.groupby("Product")["Actual_Unit_Cost_USD"].nunique() > 2).all()
    assert (t["Target_Qty"] == 0).any()                               # unplanned (new account) lines exist


def test_raw_file_content_matches_generator(pipeline, scripts):
    g = scripts["gen"]
    raw, manifest = g.corrupt(g.build_truth(config.RANDOM_SEED), config.RANDOM_SEED)
    on_disk = pd.read_csv(pipeline["manifest"], dtype=str, keep_default_na=False)
    assert len(on_disk) == len(manifest)
    assert pd.read_excel(pipeline["raw"]).shape == raw.shape


# ------------------------------------------------------------ cleaning
@pytest.fixture(scope="module")
def truth(scripts):
    return scripts["gen"].build_truth(config.RANDOM_SEED)


def test_cleaned_data_equals_ground_truth(pipeline, truth):
    clean = _read_clean(pipeline["clean_csv"])
    assert clean["Line_ID"].is_unique
    assert set(clean["Line_ID"]) <= set(truth["Line_ID"])
    m = truth.merge(clean, on="Line_ID", suffixes=("_t", "_c"), validate="one_to_one")
    for col in config.CLEAN_COLUMNS[1:]:
        a, b = m[col + "_t"], m[col + "_c"]
        if col == "Date":
            a, b = pd.to_datetime(a), pd.to_datetime(b)
        elif col == "Notes":
            a, b = a.fillna(""), b.fillna("")
        if pd.api.types.is_numeric_dtype(a):
            assert np.array_equal(np.rint(a * 100), np.rint(b * 100)), col
        else:
            assert (a.astype(str) == b.astype(str)).all(), col


def test_only_planned_review_cases_are_missing(pipeline, truth):
    clean = _read_clean(pipeline["clean_csv"])
    manifest = pd.read_csv(pipeline["manifest"], dtype=str, keep_default_na=False)
    must_review = set(manifest.loc[manifest["Expected_Handling"].str.startswith("Review"), "Line_ID"])
    missing = set(truth["Line_ID"]) - set(clean["Line_ID"])
    assert must_review, "manifest has no review cases"
    assert not (must_review & set(clean["Line_ID"])), "a row that needed review reached the clean data"
    assert missing == must_review


def test_cleaning_workbook_checks_and_reconciliation(pipeline):
    checks = pd.read_excel(pipeline["clean_xlsx"], sheet_name="Validation_Checks")
    assert len(checks) >= 15 and (checks["Result"] == "PASS").all()
    recon = pd.read_excel(pipeline["clean_xlsx"], sheet_name="Row_Reconciliation")
    rows = dict(zip(recon["Step"], recon["Rows"], strict=True))
    assert sum(v for k, v in rows.items() if not k.startswith("=")) == rows["= clean rows"]
    assert rows["= clean rows"] == len(_read_clean(pipeline["clean_csv"]))


def test_audit_lists_every_repaired_row(pipeline):
    clean = _read_clean(pipeline["clean_csv"])
    audit = pd.read_excel(pipeline["clean_xlsx"], sheet_name="Corrections_Audit", dtype={"Line_ID": str})
    repaired = set(clean.loc[clean["Repaired_Fields"].fillna("") != "", "Line_ID"])
    assert repaired <= set(audit["Line_ID"])
    assert audit["Rule"].notna().all() and (audit["Certainty"] == "Exact").all()


def test_text_starting_with_equals_is_not_a_formula(pipeline):
    ws = load_workbook(pipeline["clean_xlsx"])["Row_Reconciliation"]
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert cells and all(c.data_type == "s" for c in cells)


# ------------------------------------------------------------ analysis
def test_analysis_validation_all_pass(pipeline):
    v = pd.read_excel(pipeline["analysis_xlsx"], sheet_name="Validation_Summary")
    assert len(v) >= 30 and (v["Status"] == "PASS").all(), v[v["Status"] != "PASS"]


def test_analysis_workbook_sheets(pipeline):
    names = pd.ExcelFile(pipeline["analysis_xlsx"]).sheet_names
    for s in ["Executive_Summary", "Validation_Summary", "Metric_Dictionary", "Company", "By_Region",
              "By_Year_Month", "By_Price_Band", "Target_Accuracy", "Exceptions_RegCat_Month", "Material_Lines"]:
        assert s in names


def test_company_totals_recomputed_independently(pipeline):
    """Re-add the clean CSV text with Decimal and compare to the workbook's Company sheet."""
    raw = pd.read_csv(pipeline["clean_csv"], dtype=str)
    tot = {c: sum(Decimal(x) for x in raw[c]) for c in config.MONEY_COLUMNS if "Sales" in c or "Profit" in c}
    comp = pd.read_excel(pipeline["analysis_xlsx"], sheet_name="Company").iloc[0]
    sales_var = tot["Actual_Sales_USD"] - tot["Target_Sales_USD"]
    gp_var = tot["Actual_Gross_Profit_USD"] - tot["Target_Gross_Profit_USD"]
    assert Decimal(str(round(comp["Sales_Variance_USD"], 2))) == sales_var
    assert Decimal(str(round(comp["Gross_Profit_Variance_USD"], 2))) == gp_var
    for k in ("Volume_Pure_USD", "Mix_Effect_USD", "Price_Effect_USD"):
        assert k in comp.index
    bridge = Decimal(str(round(comp["Volume_Pure_USD"] + comp["Mix_Effect_USD"] + comp["Price_Effect_USD"], 2)))
    assert bridge == sales_var


def test_status_labels_in_outputs(pipeline):
    e = pd.read_csv(pipeline["enriched"], usecols=["Sales_Variance_Status"])
    assert set(e["Sales_Variance_Status"]) <= {config.STATUS_FAVORABLE, config.STATUS_WITHIN,
                                               config.STATUS_UNFAVORABLE, config.STATUS_NO_TARGET}
    assert "On Target" not in pipeline["html"].read_text(encoding="utf-8")


# ------------------------------------------------------- insights / README
def test_key_metrics_match_the_workbook(pipeline):
    claims = {c["key"]: c for c in json.loads(pipeline["metrics"].read_text(encoding="utf-8"))["claims"]}
    comp = pd.read_excel(pipeline["analysis_xlsx"], sheet_name="Company").iloc[0]
    pairs = {"sales_variance_usd": "Sales_Variance_USD", "sales_variance_pct": "Sales_Variance_Pct",
             "gp_variance_usd": "Gross_Profit_Variance_USD", "gp_variance_pct": "Gross_Profit_Variance_Pct",
             "price_effect_usd": "Price_Effect_USD", "pure_volume_usd": "Volume_Pure_USD", "mix_usd": "Mix_Effect_USD",
             "cost_effect_usd": "Profit_Cost_Effect_USD"}
    for key, col in pairs.items():
        assert claims[key]["expected"] == pytest.approx(float(comp[col]), abs=0.005), key
    assert len(claims) >= 20


def test_insights_report_has_no_unfilled_placeholders(pipeline):
    text = pipeline["insights"].read_text(encoding="utf-8")
    assert "nan" not in re.findall(r"\bnan\b", text.lower()) and "{" not in text and "None" not in text
    for section in ("## 1.", "## 2.", "## 3.", "R1", "R5"):
        assert section in text


def test_readme_results_block_is_generated(pipeline):
    text = pipeline["readme"].read_text(encoding="utf-8")
    start, end = text.index("<!-- KEY-RESULTS:START -->"), text.index("<!-- KEY-RESULTS:END -->")
    assert start < end and "$" in text[start:end]


# --------------------------------------------- committed files are current
@pytest.mark.parametrize("key, repo_path", [
    ("clean_csv", config.CLEAN_CSV), ("enriched", config.ENRICHED_CSV), ("manifest", config.DEFECT_MANIFEST_FILE),
    ("insights", config.INSIGHTS_MD), ("metrics", config.KEY_METRICS_JSON), ("readme", ROOT / "README.md"),
])
def test_committed_outputs_are_up_to_date(pipeline, key, repo_path):
    """Fails if code changed but the committed outputs were not regenerated."""
    assert repo_path.is_file(), f"{repo_path} missing - run the pipeline"
    fresh = pipeline[key].read_bytes().replace(b"\r\n", b"\n")
    committed = repo_path.read_bytes().replace(b"\r\n", b"\n")
    assert fresh == committed, f"{repo_path.relative_to(ROOT)} is stale - rerun the pipeline and commit"


# ------------------------------------------------------------ dashboard
def test_dashboard_is_self_contained(pipeline):
    html = pipeline["html"].read_text(encoding="utf-8")
    assert not re.search(r"<script[^>]+src=", html, re.I)
    assert not re.search(r"<link[^>]+stylesheet", html, re.I)
    assert html.lower().count("</script") == html.lower().count("<script")   # data cannot close the script tag
    assert "On Target" not in html


# ------------------------------------------------------ failure paths
def test_analysis_rejects_missing_input(scripts, tmp_path, capsys):
    assert scripts["analysis"].main(["--input", str(tmp_path / "nope.csv"), "--output-dir", str(tmp_path),
                                     "--workbook-dir", str(tmp_path)]) == 2
    assert "ERROR" in capsys.readouterr().err


def test_insights_rejects_missing_input(scripts, tmp_path):
    assert scripts["insights"].main(["--enriched", str(tmp_path / "x.csv"), "--workbook", str(tmp_path / "x.xlsx"),
                                     "--docs-dir", str(tmp_path), "--readme", str(tmp_path / "R.md")]) == 2


def test_dashboard_rejects_missing_input(scripts, tmp_path):
    assert scripts["dashboard"].main(["--enriched", str(tmp_path / "x.csv"), "--output", str(tmp_path / "d.html")]) != 0
    assert not (tmp_path / "d.html").exists()
